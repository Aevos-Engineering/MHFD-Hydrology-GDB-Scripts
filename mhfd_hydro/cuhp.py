"""Read subcatchment inputs from a CUHP 2005 (v2.x) workbook (.xlsm/.xlsx)."""

from dataclasses import dataclass

# Column order of the CUHP 2005 v2.0.x "Subcatchments" sheet.
_COLS = ["name", "target_node", "raingage", "area_sq_mi", "length_to_centroid_mi",
         "length_mi", "slope_ftft", "impervious_pct", "depression_pervious_in",
         "depression_impervious_in", "infil_initial_inhr", "infil_decay_1sec",
         "infil_final_inhr", "dcia_level"]
_OVERRIDE_COLS = [14, 16, 19, 21, 23, 24, 25, 26]


@dataclass
class Subcatchment:
    name: str
    target_node: str
    raingage: str
    area_sq_mi: float
    length_to_centroid_mi: float
    length_mi: float
    slope_ftft: float
    impervious_pct: float
    depression_pervious_in: float
    depression_impervious_in: float
    infil_initial_inhr: float
    infil_decay_1sec: float
    infil_final_inhr: float
    dcia_level: float
    has_override: bool = False
    impervious_existing_pct: float = None
    impervious_future_pct: float = None

    @property
    def area_ac(self):
        return None if self.area_sq_mi is None else self.area_sq_mi * 640.0


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def read_cuhp(path):
    """Return (subcatchments list, info dict) from a CUHP workbook."""
    import openpyxl
    import warnings

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    info = {"path": str(path)}
    if "Intro" in wb.sheetnames:
        for row in wb["Intro"].iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip().lower().startswith("version"):
                    info.setdefault("version", v.strip())
            vals = [v for v in row if v not in (None, "")]
            if len(vals) >= 2 and isinstance(vals[0], str):
                key = vals[0].strip().rstrip(":").lower()
                if key in ("project title", "project comment"):
                    info[key.replace(" ", "_")] = vals[1]

    ws = wb["Subcatchments"]
    subs, started = [], False
    for row in ws.iter_rows(values_only=True):
        if not started:
            if row and isinstance(row[0], str) and row[0].strip().lower() == "subcatchment name":
                started = True
            continue
        if not row or row[0] in (None, ""):
            continue
        vals = dict(zip(_COLS, row[:len(_COLS)]))
        sc = Subcatchment(
            name=str(vals["name"]).strip(),
            target_node=str(vals["target_node"]).strip() if vals["target_node"] is not None else None,
            raingage=vals["raingage"],
            **{k: _num(vals[k]) for k in _COLS[3:]},
        )
        sc.has_override = any(len(row) > c and row[c] not in (None, "") for c in _OVERRIDE_COLS)
        subs.append(sc)

    # Existing/future imperviousness, when the Multiple Runs sheet is filled in.
    if "Multiple Runs" in wb.sheetnames:
        by_name = {s.name: s for s in subs}
        hdr = False
        for row in wb["Multiple Runs"].iter_rows(values_only=True):
            if not hdr:
                hdr = bool(row and isinstance(row[0], str) and row[0].strip().lower() == "subcatchment name")
                continue
            if row and row[0] in by_name:
                by_name[row[0]].impervious_existing_pct = _num(row[1])
                by_name[row[0]].impervious_future_pct = _num(row[2])
    info["rainfall"] = read_rainfall(wb)
    return subs, info


def read_rainfall(wb):
    """1-hr and 6-hr depths from each CUHP storm sheet (rows '1Hr Depth',
    '6Hr Depth', 'Return Period'), one entry per return period.

    Returns a list of {"return_period", "one_hour_in", "six_hour_in", "sheet"}.
    """
    from . import constants as C

    out, seen = [], set()
    for ws in wb.worksheets:
        vals = {}
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True)):
            cells = list(row or ())
            for j, v in enumerate(cells[:-1]):
                if isinstance(v, str):
                    k = v.strip().lower().replace(" ", "")
                    if k in ("1hrdepth", "6hrdepth", "returnperiod"):
                        vals.setdefault(k, cells[j + 1])
        one, six, rp = _num(vals.get("1hrdepth")), _num(vals.get("6hrdepth")), vals.get("returnperiod")
        if one is None or rp is None:
            continue
        rp_n = C.normalize_return_period(int(rp) if isinstance(rp, float) and rp.is_integer() else rp)
        if not rp_n or rp_n in seen:
            continue
        seen.add(rp_n)
        out.append({"return_period": rp_n, "one_hour_in": one, "six_hour_in": six, "sheet": ws.title})
    order = {k: i for i, k in enumerate(C.RETURN_PERIODS)}
    return sorted(out, key=lambda r: order.get(r["return_period"], 99))
